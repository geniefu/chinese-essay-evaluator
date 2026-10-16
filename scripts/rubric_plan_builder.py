# -*- coding: utf-8 -*-
"""
寫作測驗嚴格評審準則規畫書 Word (.docx) 自動生成器 (rubric_plan_builder.py)
用途：
接收作文題目、引導材料與評審要求，自動產出具備專業排版的《[題目]_嚴格評審準則規畫書.docx》，
供老師審查、微調，並作為後續批量批閱之嚴格評分規準 (Ground Truth)。
"""

import os
import sys
import argparse
import json
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

def set_font(run, name='微軟正黑體', size_pt=10.0, bold=False, color_rgb=(15, 23, 42)):
    """設定文字樣式 (中文字體微軟正黑體，英數字 Times New Roman)"""
    run.font.name = 'Times New Roman'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Times New Roman')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color_rgb)

def set_cell_shd(cell, color_hex):
    """設定表格單元格背景色"""
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_pad(cell, top=60, bottom=60, left=90, right=90):
    """設定單元格內邊距"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_borders(cell, top=None, bottom=None, left=None, right=None):
    """設定單元格框線"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="{top.get('val', 'none') if top else 'none'}" w:sz="{top.get('sz', '4') if top else '0'}" w:space="0" w:color="{top.get('color', 'auto') if top else 'auto'}"/>
            <w:left w:val="{left.get('val', 'none') if left else 'none'}" w:sz="{left.get('sz', '4') if left else '0'}" w:space="0" w:color="{left.get('color', 'auto') if left else 'auto'}"/>
            <w:bottom w:val="{bottom.get('val', 'none') if bottom else 'none'}" w:sz="{bottom.get('sz', '4') if bottom else '0'}" w:space="0" w:color="{bottom.get('color', 'auto') if bottom else 'auto'}"/>
            <w:right w:val="{right.get('val', 'none') if right else 'none'}" w:sz="{right.get('sz', '4') if right else '0'}" w:space="0" w:color="{right.get('color', 'auto') if right else 'auto'}"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)

def add_callout(doc, text, title="重點提示", border_color="1E3A8A", bg_color="EFF6FF"):
    """加入精緻 Callout 提示框"""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(7.1)
    set_cell_shd(cell, bg_color)
    set_cell_pad(cell, top=100, bottom=100, left=140, right=140)
    set_borders(cell, left={'val': 'single', 'sz': '24', 'color': border_color},
                      top={'val': 'none'}, bottom={'val': 'none'}, right={'val': 'none'})
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.2
    r_t = p.add_run(f"【{title}】\n")
    c_rgb = (30, 58, 138) if border_color=="1E3A8A" else ((180, 83, 9) if border_color=="D97706" else (185, 28, 28))
    set_font(r_t, name='微軟正黑體', size_pt=10.0, bold=True, color_rgb=c_rgb)
    r_c = p.add_run(text)
    set_font(r_c, name='微軟正黑體', size_pt=9.5, bold=False, color_rgb=(51, 65, 85))

def format_heading(doc, text, level=1):
    """加入章節標題"""
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    if level == 1:
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(text)
        set_font(run, name='微軟正黑體', size_pt=12.5, bold=True, color_rgb=(30, 58, 138))
    elif level == 2:
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(text)
        set_font(run, name='微軟正黑體', size_pt=10.5, bold=True, color_rgb=(37, 99, 235))
    return p

def add_p(doc, text="", space_after=3, bold=False, color_rgb=(30, 41, 59)):
    """加入正文段落"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.2
    if text:
        run = p.add_run(text)
        set_font(run, name='微軟正黑體', size_pt=9.5, bold=bold, color_rgb=color_rgb)
    return p

def generate_rubric_plan_docx(plan_data, output_path):
    """根據規畫書資料結構生成 Word 檔案"""
    doc = docx.Document()

    # 設定標準邊界 (0.75 吋 / 約 1.9 公分)
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    topic = plan_data.get('topic', '未定題目')
    stage = plan_data.get('stage', '國中九年級（會考衝刺）')
    mode = plan_data.get('mode', '會考高標閱卷／從嚴評定模式')
    guidance = plan_data.get('guidance', '請依引導寫作說明撰寫。')
    core_theme = plan_data.get('core_theme', '')
    boundaries = plan_data.get('boundaries', {})
    rubric_matrix = plan_data.get('rubric_matrix', [])
    strict_traps = plan_data.get('strict_traps', [])
    model_essay_outline = plan_data.get('model_essay_outline', {})
    checklist = plan_data.get('checklist', [])

    # 1. 主標題
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    r_title = p_title.add_run(f"〈{topic}〉寫作測驗嚴格評審準則規畫書")
    set_font(r_title, name='微軟正黑體', size_pt=16, bold=True, color_rgb=(30, 58, 138))

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(8)
    r_sub = p_sub.add_run(f"【適用學段：{stage} ｜ 評審尺度：{mode}】")
    set_font(r_sub, name='微軟正黑體', size_pt=10, bold=False, color_rgb=(100, 116, 139))

    # 2. 命題資訊橫卡 (2x2 表格)
    tbl_info = doc.add_table(rows=2, cols=2)
    tbl_info.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_info.autofit = False
    info_widths = [Inches(3.5), Inches(3.6)]
    card_items = [
        [("寫作題目", f"〈{topic}〉"), ("適用對象與標準", stage)],
        [("閱卷標準模式", mode), ("預定輸出規範", "雙檔案 (Word + PDF) 診斷單與 Excel 總表")]
    ]
    for r_idx in range(2):
        for c_idx in range(2):
            cell = tbl_info.cell(r_idx, c_idx)
            cell.width = info_widths[c_idx]
            lbl, val = card_items[r_idx][c_idx]
            set_cell_shd(cell, "F8FAFC")
            set_cell_pad(cell, top=60, bottom=60, left=80, right=80)
            set_borders(cell, top={'val': 'single', 'sz': '4', 'color': 'E2E8F0'},
                              bottom={'val': 'single', 'sz': '4', 'color': 'E2E8F0'},
                              left={'val': 'single', 'sz': '4', 'color': 'E2E8F0'},
                              right={'val': 'single', 'sz': '4', 'color': 'E2E8F0'})
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
            r1 = p.add_run(f"{lbl}：")
            set_font(r1, name='微軟正黑體', size_pt=9.0, bold=True, color_rgb=(71, 85, 105))
            r2 = p.add_run(val)
            set_font(r2, name='微軟正黑體', size_pt=9.0, bold=False, color_rgb=(15, 23, 42))

    # 3. 命題引導材料 Callout
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    add_callout(doc, guidance, title="📋 命題引導寫作材料與限制條件", border_color="1E3A8A", bg_color="EFF6FF")

    # 4. 壹、排他性審題拆解與取材邊界
    format_heading(doc, "壹、 排他性審題拆解與取材邊界", level=1)
    if core_theme:
        add_p(doc, f"● 核心命題立意：{core_theme}", bold=True, color_rgb=(30, 58, 138))

    tbl_bound = doc.add_table(rows=4, cols=2)
    tbl_bound.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_bound.autofit = False
    bound_rows = [
        ("✅ 核心切題範圍（優等）", boundaries.get('in_scope', '能精準切入題旨，結合真實生活經驗，有具體物象描摹與深度反思。')),
        ("⚠️ 踩線／擦邊選材（扣分預警）", boundaries.get('borderline', '選材流於泛泛之論或生活瑣事流水帳，缺乏明確聚焦點，最高鎖定 4 級分。')),
        ("❌ 偏題／違背限制（嚴格降級）", boundaries.get('off_scope', '脫離引導情境限制、未符題目關鍵字核心義，最高鎖定 2~3 級分。')),
        ("🚫 嚴重違規（0 級分判定）", boundaries.get('zero_tolerance', '空白卷、完全抄題、詩歌體裁、洩漏個人身分資料。'))
    ]
    for idx, (title_b, desc_b) in enumerate(bound_rows):
        c_title = tbl_bound.cell(idx, 0); c_title.width = Inches(2.3)
        c_desc = tbl_bound.cell(idx, 1); c_desc.width = Inches(4.8)
        shd_col = "F1F5F9" if idx == 0 else ("FEF3C7" if idx == 1 else ("FEE2E2" if idx == 2 else "F3F4F6"))
        border_col = "CBD5E1" if idx == 0 else ("F59E0B" if idx == 1 else ("EF4444" if idx == 2 else "9CA3AF"))
        for cell in (c_title, c_desc):
            set_cell_pad(cell, top=60, bottom=60, left=80, right=80)
            set_borders(cell, top={'val': 'single', 'sz': '4', 'color': border_col},
                              bottom={'val': 'single', 'sz': '4', 'color': border_col},
                              left={'val': 'single', 'sz': '4', 'color': border_col},
                              right={'val': 'single', 'sz': '4', 'color': border_col})
        set_cell_shd(c_title, shd_col)
        set_cell_shd(c_desc, "FFFFFF")
        
        p_t = c_title.paragraphs[0]; p_t.paragraph_format.space_before = Pt(0); p_t.paragraph_format.space_after = Pt(0)
        r_t = p_t.add_run(title_b)
        set_font(r_t, name='微軟正黑體', size_pt=9.0, bold=True, color_rgb=(30, 41, 59))

        p_d = c_desc.paragraphs[0]; p_d.paragraph_format.space_before = Pt(0); p_d.paragraph_format.space_after = Pt(0)
        p_d.paragraph_format.line_spacing = 1.15
        r_d = p_d.add_run(desc_b)
        set_font(r_d, name='微軟正黑體', size_pt=9.0, bold=False, color_rgb=(51, 65, 85))

    # 5. 貳、嚴格閱卷 0～6 級分量化指標對照表
    format_heading(doc, "貳、 嚴格閱卷 0～6 級分量化指標對照表", level=1)
    
    headers = ["綜合級分", "得分核心特徵", "扣分地雷與上限鎖定", "本題具體表現判準"]
    widths = [Inches(1.0), Inches(2.2), Inches(2.0), Inches(1.9)]
    
    # 預設 6~0 級分標準矩陣（若使用者未傳入特定矩陣時自適應填充）
    default_matrix = [
        {"score": "6 級分", "trait": "篇幅充裕跨頁；審題透徹；雙事例以上層次遞進；由物入理哲思昇華；極少錯字。", "trap": "無客觀畫面或只抒情者不得列入；單一事例降級。", "spec": "能從本題情境深刻提煉出生命洞見，文字精準富美感。"},
        {"score": "5 級分", "trait": "篇幅充裕達標；切題具體生動；事例充足；文句流暢有變化；少有錯字。", "trap": "選材略顯常見但敘事真摯者上限 5 級分；事例不足 2 例者降級。", "spec": "清楚扣合引導說明，情感真誠，能由事至理完整開展。"},
        {"score": "4 級分", "trait": "基本切題，結構完整。取材停留在表面，感悟略顯口號化。", "trap": "① 內容略少（未滿一頁且行數未達90%）<br>② 俗套選材無新意<br>③ 單一事例單薄", "spec": "平鋪直敘完成題目要求，符合大眾常規構思，深度有限。"},
        {"score": "3 級分", "trait": "嘗試切題但開展不足；文筆平淡；流水帳或文句略有阻滯。", "trap": "① 篇幅太少（未滿一頁且行數未達75%）<br>② 通篇抽象雞湯無具體畫面<br>③ 偏離引導限制", "spec": "篇幅明顯不足，或舉例不合題意，強制觸發寫作進階鷹架。"},
        {"score": "2 級分", "trait": "嚴重偏題；材料貧乏破碎；行數嚴重短缺；錯字連篇。", "trap": "未入題、不知所云、文不成段。", "spec": "僅寫數行或前後文意毫無關聯。"},
        {"score": "1 級分", "trait": "嚴重離題，支離破碎，僅寥寥數語。", "trap": "極度不足、字跡無法辨識。", "spec": "完全無法成篇。"},
        {"score": "0 級分", "trait": "空白卷、完全抄題、詩歌體、洩漏個人身分。", "trap": "考場違規條款。", "spec": "依大考規章給予零分。"}
    ]
    matrix_data = rubric_matrix if rubric_matrix else default_matrix

    tbl_matrix = doc.add_table(rows=len(matrix_data) + 1, cols=4)
    tbl_matrix.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_matrix.autofit = False

    # 表頭
    for idx_c, h in enumerate(headers):
        c = tbl_matrix.cell(0, idx_c); c.width = widths[idx_c]
        set_cell_shd(c, "1E3A8A")
        set_cell_pad(c, top=60, bottom=60, left=60, right=60)
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
        run = p.add_run(h)
        set_font(run, name='微軟正黑體', size_pt=9.0, bold=True, color_rgb=(255, 255, 255))

    # 資料列
    for r_idx, row in enumerate(matrix_data):
        sc_cell = tbl_matrix.cell(r_idx + 1, 0); sc_cell.width = widths[0]
        tr_cell = tbl_matrix.cell(r_idx + 1, 1); tr_cell.width = widths[1]
        tp_cell = tbl_matrix.cell(r_idx + 1, 2); tp_cell.width = widths[2]
        sp_cell = tbl_matrix.cell(r_idx + 1, 3); sp_cell.width = widths[3]

        bg_c = "F8FAFC" if r_idx % 2 == 0 else "FFFFFF"
        if "6" in str(row['score']):
            bg_c = "EFF6FF"
        elif "4" in str(row['score']):
            bg_c = "FEFCE8"
        elif "3" in str(row['score']) or "2" in str(row['score']):
            bg_c = "FEF2F2"

        for cell in (sc_cell, tr_cell, tp_cell, sp_cell):
            set_cell_pad(cell, top=50, bottom=50, left=60, right=60)
            set_cell_shd(cell, bg_c)
            set_borders(cell, top={'val': 'single', 'sz': '4', 'color': 'E2E8F0'},
                              bottom={'val': 'single', 'sz': '4', 'color': 'E2E8F0'},
                              left={'val': 'single', 'sz': '4', 'color': 'E2E8F0'},
                              right={'val': 'single', 'sz': '4', 'color': 'E2E8F0'})

        p_sc = sc_cell.paragraphs[0]; p_sc.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_sc.paragraph_format.space_before = Pt(0); p_sc.paragraph_format.space_after = Pt(0)
        r_sc = p_sc.add_run(str(row['score']))
        set_font(r_sc, name='微軟正黑體', size_pt=9.0, bold=True, color_rgb=(30, 58, 138) if "6" in str(row['score']) else (15, 23, 42))

        for cell_obj, key in ((tr_cell, 'trait'), (tp_cell, 'trap'), (sp_cell, 'spec')):
            p = cell_obj.paragraphs[0]; p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            txt = row.get(key, '').replace('<br>', '\n')
            r = p.add_run(txt)
            set_font(r, name='微軟正黑體', size_pt=8.5, bold=False, color_rgb=(51, 65, 85))

    # 6. 參、高標閱卷扣分地雷與防守線細則
    format_heading(doc, "參、 高標閱卷扣分地雷與防守線細則", level=1)
    traps = strict_traps if strict_traps else [
        "【篇幅防守線】未滿一頁且行數未達 75% 強制鎖定 3 級分以下，並啟動寫作鷹架；未達 90% 最高鎖定 4 級分。",
        "【俗套題材限制】若取材於生病發燒、段考失利、出遊塞車等常見俗套，且未能提供嶄新哲思視角者，立意取材最高鎖定 4 級分。",
        "【單一事例限制】通篇僅舉一個生活事件者，論證層次不足，上限鎖定 4 級分；欲衝刺 5～6 級分必須有 2 個以上事例層次遞進。",
        "【客觀畫面檢核】通篇僅有抽象抒情、雞湯口號而無具體特寫畫面（視覺/聽覺等）者，不得評入 5 級分以上。",
        "【錯字與座標定位】錯別字一律還原考場天頭座標定位（P1 第 X 行第 Y 格）；錯字達 3~5 個以上者在級分邊界從嚴降級。"
    ]
    for trap_item in traps:
        add_p(doc, f"● {trap_item}", space_after=3, color_rgb=(185, 28, 28))

    # 7. 肆、6 級分標竿範文架構與取材走向
    format_heading(doc, "肆、 6 級分標竿範文架構與取材走向", level=1)
    outline = model_essay_outline if model_essay_outline else {
        "起（破題扣合）": "以具體客觀事物或生活微觀場景破題，自然帶出核心主題，絕不以流水帳入題。",
        "承（細節開展）": "精雕細琢第一事例之客觀細節，融入五感描摹（光影、聲音、觸感），展現真實生活感。",
        "轉（深化遞進）": "引出第二事例或心境轉折，由事入理，跳脫個人狹隘視角，連結更廣闊的人生體驗。",
        "合（哲思收束）": "首尾呼應，從具體生活觀察昇華為普遍性生命智慧，餘音繞樑。"
    }
    for seg_title, seg_desc in outline.items():
        p_seg = add_p(doc, f"【{seg_title}】{seg_desc}", space_after=2)

    # 8. 伍、批閱執行核准與簽章確認欄
    format_heading(doc, "伍、 評審準則審查與確認（人機核准欄）", level=1)
    add_callout(doc, 
        "本規畫書產出後，請教師親自檢閱審題條件與分級判準：\n"
        "1. 若同意本規畫書內容：請回覆「確認同意，請依此規畫書批改指定資料夾」即可自動啟動批量批閱。\n"
        "2. 若有增修項目：可直接修改本 Word 檔案儲存後指向該檔案，或直接於對話中提出修正指示重新生成。",
        title="📝 教師審查指示", border_color="D97706", bg_color="FFFBEB")

    # 存檔
    doc.save(output_path)
    print(f"[OK] 嚴格評審準則規畫書已成功生成：{output_path}")

def main():
    parser = argparse.ArgumentParser(description="嚴格評審準則規畫書 Word (.docx) 生成器")
    parser.add_argument('--topic', required=True, help="作文題目")
    parser.add_argument('--stage', default="國中九年級（會考衝刺）", help="適用學段與年級")
    parser.add_argument('--mode', default="會考高標閱卷／從嚴評定模式", help="評審尺度模式")
    parser.add_argument('--guidance', default="請依命題引導材料撰寫。", help="引導寫作說明")
    parser.add_argument('--core-theme', default="", help="核心立意宗旨")
    parser.add_argument('--output', required=True, help="輸出 .docx 檔案路徑")
    parser.add_argument('--data-json', help="傳入完整 JSON 設定檔路徑")

    args = parser.parse_args()

    plan_data = {}
    if args.data_json and os.path.exists(args.data_json):
        with open(args.data_json, 'r', encoding='utf-8') as f:
            plan_data = json.load(f)
    
    if args.topic:
        plan_data['topic'] = args.topic
    if args.stage:
        plan_data['stage'] = args.stage
    if args.mode:
        plan_data['mode'] = args.mode
    if args.guidance:
        plan_data['guidance'] = args.guidance
    if args.core_theme:
        plan_data['core_theme'] = args.core_theme

    generate_rubric_plan_docx(plan_data, args.output)

if __name__ == '__main__':
    main()
