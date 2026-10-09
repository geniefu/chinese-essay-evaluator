# -*- coding: utf-8 -*-
"""
寫作測驗專業評閱診斷單 Word (.docx) 自動生成器 (essay_report_builder.py)
支援模式：
1. 單生專屬診斷卡 (Single Report)
2. 全班成果彙整手冊 (Class-wide Batch Report)：含封面總覽表、每生強制分頁獨立專章、全班教學診斷專頁
"""

import os
import sys
import argparse
import json
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

def set_font(run, name='微軟正黑體', size_pt=10.0, bold=False, color_rgb=(15, 23, 42)):
    """設定文字樣式 (中文字體預設微軟正黑體/標楷體，英數字 Times New Roman)"""
    run.font.name = 'Times New Roman'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Times New Roman')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color_rgb)

def set_cell_shd(cell, color_hex):
    """設定表格單元格背景色"""
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_pad(cell, top=50, bottom=50, left=70, right=70):
    """設定單元格內邊距"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_borders(cell, top=None, bottom=None, left=None, right=None):
    """設定單元格框線"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="{top.get('val', 'none') if top else 'none'}" w:sz="{top.get('sz', '4') if top else '0'}" w:space="0" w:color="{top.get('color', 'auto') if top else 'auto'}"/>
            <w:left w:val="{left.get('val', 'none') if left else 'none'}" w:sz="{left.get('sz', '4') if left else '0'}" w:space="0" w:color="{left.get('color', 'auto') if left else 'auto'}"/>
            <w:bottom w:val="{bottom.get('val', 'none') if bottom else 'none'}" w:sz="{bottom.get('sz', '4') if bottom else '0'}" w:space="0" w:color="{bottom.get('color', 'auto') if bottom else 'auto'}"/>
            <w:right w:val="{right.get('val', 'none') if right else 'none'}" w:sz="{right.get('sz', '4') if right else '0'}" w:space="0" w:color="{right.get('color', 'auto') if right else 'auto'}"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)

def add_callout(doc, text, title="重點提示", border_color="1E3A8A", bg_color="EFF6FF"):
    """加入精緻 Callout 提示框"""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(7.1)
    set_cell_shd(cell, bg_color)
    set_cell_pad(cell, top=100, bottom=100, left=140, right=140)
    set_borders(cell, left={'val': 'single', 'sz': '24', 'color': border_color},
                      top={'val': 'none'}, bottom={'val': 'none'}, right={'val': 'none'})
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.2
    r_t = p.add_run(f"【{title}】\n")
    c_rgb = (30, 58, 138) if border_color=="1E3A8A" else ((180, 83, 9) if border_color=="D97706" else (185, 28, 28))
    set_font(r_t, name='微軟正黑體', size_pt=10.0, bold=True, color_rgb=c_rgb)
    r_c = p.add_run(text)
    set_font(r_c, name='微軟正黑體', size_pt=9.0, bold=False, color_rgb=(51, 65, 85))

def format_heading(doc, text, level=1):
    """加入章節標題"""
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    if level == 1:
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(text)
        set_font(run, name='微軟正黑體', size_pt=12, bold=True, color_rgb=(30, 58, 138))
    elif level == 2:
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(text)
        set_font(run, name='微軟正黑體', size_pt=10.5, bold=True, color_rgb=(37, 99, 235))
    return p

def add_p(doc, text="", space_after=2, bold=False, color_rgb=(30, 41, 59)):
    """加入正文段落"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if text:
        run = p.add_run(text)
        set_font(run, name='微軟正黑體', size_pt=9.0, bold=bold, color_rgb=color_rgb)
    return p

def render_student_section(doc, student_data, is_batch=False):
    """渲染單一學生評閱章節"""
    topic = student_data.get('topic', "未定題")
    student_info = student_data.get('student_info', "九年級 學生")
    score = student_data.get('score', 4)
    percent_score = student_data.get('percent_score', 80)
    level_tier = student_data.get('tier', "中卷穩健 (B++)")
    stage = student_data.get('stage', "國中九年級")
    one_sentence_review = student_data.get('one_sentence_review', "符合題意，但取材多停留在表面觀察，感悟流於口號化。")
    page_status = student_data.get('page_status', "卷面行數約 85%（最高上限鎖定 4 級分）")
    dimensions = student_data.get('dimensions', [])
    wrong_words = student_data.get('wrong_words', [])
    section_comments = student_data.get('section_comments', [])
    upgrade_demo = student_data.get('upgrade_demo', {})
    scaffold = student_data.get('scaffold', None)
    model_essay = student_data.get('model_essay', None)

    # 1. 學生章節標題
    p_header = doc.add_paragraph()
    p_header.paragraph_format.space_before = Pt(4)
    p_header.paragraph_format.space_after = Pt(3)
    p_header.paragraph_format.keep_with_next = True
    r_hdr = p_header.add_run(f"🧑‍🎓 {student_info} 寫作測驗專業評閱診斷報告")
    set_font(r_hdr, name='微軟正黑體', size_pt=14, bold=True, color_rgb=(30, 58, 138))

    # 2. 基本資訊橫卡
    tbl_info = doc.add_table(rows=1, cols=4)
    tbl_info.alignment = WD_TABLE_ALIGNMENT.CENTER; tbl_info.autofit = False
    info_w = [Inches(2.5), Inches(1.6), Inches(1.6), Inches(1.5)]
    info_items = [
        ("題目", f"〈{topic}〉"),
        ("學段/對象", f"{stage} ｜ {student_info}"),
        ("評定級分", f"{score} 級分 ({level_tier})"),
        ("百分制換算", f"{percent_score} 分")
    ]
    for idx_i, (lbl, val) in enumerate(info_items):
        c = tbl_info.cell(0, idx_i); c.width = info_w[idx_i]
        set_cell_shd(c, "F1F5F9")
        set_cell_pad(c, top=40, bottom=40, left=50, right=50)
        set_borders(c, top={'val': 'single', 'sz': '4', 'color': 'CBD5E1'}, bottom={'val': 'single', 'sz': '4', 'color': 'CBD5E1'})
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r1 = p.add_run(f"{lbl}\n"); set_font(r1, name='微軟正黑體', size_pt=8.0, bold=False, color_rgb=(100, 116, 139))
        r2 = p.add_run(val); set_font(r2, name='微軟正黑體', size_pt=9.0, bold=True, color_rgb=(30, 58, 138))

    add_p(doc, f"卷面篇幅與防守線檢核：{page_status}", space_after=3, bold=True, color_rgb=(71, 85, 105))

    # 3. 老師考場紅筆總評 Banner
    tbl_score = doc.add_table(rows=1, cols=2)
    tbl_score.alignment = WD_TABLE_ALIGNMENT.CENTER; tbl_score.autofit = False
    c_sc1 = tbl_score.cell(0, 0); c_sc1.width = Inches(2.0)
    c_sc2 = tbl_score.cell(0, 1); c_sc2.width = Inches(5.2)

    sc_bg = "EFF6FF" if score >= 5 else ("FEF3C7" if score == 4 else "FEF2F2")
    sc_border = "1D4ED8" if score >= 5 else ("D97706" if score == 4 else "DC2626")
    set_cell_shd(c_sc1, sc_bg); set_cell_pad(c_sc1, top=50, bottom=50, left=30, right=30)
    set_borders(c_sc1, left={'val': 'single', 'sz': '20', 'color': sc_border},
                       top={'val': 'single', 'sz': '4', 'color': 'CBD5E1'},
                       bottom={'val': 'single', 'sz': '4', 'color': 'CBD5E1'})
    p_sc = c_sc1.paragraphs[0]; p_sc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sc_lbl = p_sc.add_run("綜合會考級分\n"); set_font(r_sc_lbl, name='微軟正黑體', size_pt=8.5, bold=False, color_rgb=(100, 116, 139))
    r_sc_val = p_sc.add_run(f"【 {score} 級分 】\n")
    set_font(r_sc_val, name='微軟正黑體', size_pt=18, bold=True, color_rgb=(30, 58, 138) if score>=5 else ((180, 83, 9) if score==4 else (185, 28, 28)))
    r_sc_tier = p_sc.add_run(level_tier); set_font(r_sc_tier, name='微軟正黑體', size_pt=8.5, bold=True, color_rgb=(71, 85, 105))

    set_cell_shd(c_sc2, "FFFFFF"); set_cell_pad(c_sc2, top=50, bottom=50, left=50, right=50)
    set_borders(c_sc2, right={'val': 'single', 'sz': '4', 'color': 'CBD5E1'},
                       top={'val': 'single', 'sz': '4', 'color': 'CBD5E1'},
                       bottom={'val': 'single', 'sz': '4', 'color': 'CBD5E1'})
    p_rev = c_sc2.paragraphs[0]
    r_rev_t = p_rev.add_run("閱卷名師實戰紅筆總評：\n"); set_font(r_rev_t, name='微軟正黑體', size_pt=9.5, bold=True, color_rgb=(30, 58, 138))
    r_rev_c = p_rev.add_run(one_sentence_review); set_font(r_rev_c, name='微軟正黑體', size_pt=9.0, bold=False, color_rgb=(51, 65, 85))

    # 4. 四大向度評量矩陣表
    format_heading(doc, "壹、 四大評分向度量化考核", 1)
    tbl_dim = doc.add_table(rows=1, cols=3)
    tbl_dim.alignment = WD_TABLE_ALIGNMENT.CENTER; tbl_dim.autofit = False
    dim_w = [Inches(1.4), Inches(1.1), Inches(4.7)]
    for i, h in enumerate(["評審向度", "向度等第", "核心診斷亮點與硬傷剖析"]):
        c = tbl_dim.rows[0].cells[i]; c.text = h; c.width = dim_w[i]
        set_cell_shd(c, "1E3A8A"); set_cell_pad(c, top=50, bottom=50, left=40, right=40)
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.runs[0], name='微軟正黑體', size_pt=8.5, bold=True, color_rgb=(255, 255, 255))

    default_dims = [
        ("立意取材", f"{min(score, 4)} 級分", "取材切合題意，但多停留在表面客觀描述，缺乏獨特生命發現與深層哲理體悟。"),
        ("結構組織", f"{score} 級分", "結構大致完整，但僅舉單一事例略顯單薄，建議落實『一例一段、2~3個事例』層次展開。"),
        ("遣詞造句", f"{score} 級分", "文句通順清楚，能傳達意思，但句型稍平鋪，偶有口語化贅字。"),
        ("錯別字與標點", f"{max(score, 4)} 級分", "格式大致正確，有少數錯別字，標點符號請注意逗號與句號之適當分界。")
    ]
    target_dims = dimensions if dimensions else default_dims

    for r_i, d_row in enumerate(target_dims):
        row = tbl_dim.add_row()
        bg = "F8FAFC" if r_i % 2 == 1 else "FFFFFF"
        for c_i, val in enumerate(d_row):
            c = row.cells[c_i]; c.text = val; c.width = dim_w[c_i]
            set_cell_shd(c, bg); set_cell_pad(c, top=45, bottom=45, left=45, right=45)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_i in [0, 1] else WD_ALIGN_PARAGRAPH.LEFT
            f_color = (30, 58, 138) if c_i == 0 else ((180, 83, 9) if c_i == 1 else (15, 23, 42))
            set_font(p.runs[0], name='微軟正黑體', size_pt=8.5, bold=(c_i in [0, 1]), color_rgb=f_color)
            set_borders(c, bottom={'val': 'single', 'sz': '4', 'color': 'CBD5E1'}, top={'val': 'single', 'sz': '4', 'color': 'CBD5E1'})

    # 5. 考場天頭錯別字訂正格
    format_heading(doc, "貳、 考場天頭錯別字回寫訂正格", 1)
    if wrong_words:
        tbl_ww = doc.add_table(rows=1, cols=4)
        tbl_ww.alignment = WD_TABLE_ALIGNMENT.CENTER; tbl_ww.autofit = False
        ww_w = [Inches(1.8), Inches(1.8), Inches(1.8), Inches(1.8)]
        for i, h in enumerate(["錯字回寫欄 ①", "錯字回寫欄 ②", "錯字回寫欄 ③", "錯字回寫欄 ④"]):
            c = tbl_ww.rows[0].cells[i]; c.text = h; c.width = ww_w[i]
            set_cell_shd(c, "334155"); set_cell_pad(c, top=45, bottom=45, left=35, right=35)
            p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_font(p.runs[0], name='微軟正黑體', size_pt=8.5, bold=True, color_rgb=(255, 255, 255))
        
        row_w = tbl_ww.add_row()
        for idx_w in range(4):
            c = row_w.cells[idx_w]; c.width = ww_w[idx_w]
            set_cell_shd(c, "FEF2F2"); set_cell_pad(c, top=50, bottom=50, left=35, right=35)
            p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if idx_w < len(wrong_words):
                w_orig, w_corr = wrong_words[idx_w]
                r_txt = p.add_run(f"「{w_orig}」 口 ➔ 【 {w_corr} 】")
                set_font(r_txt, name='微軟正黑體', size_pt=8.5, bold=True, color_rgb=(185, 28, 28))
            else:
                r_txt = p.add_run("---")
                set_font(r_txt, name='微軟正黑體', size_pt=8.5, bold=False, color_rgb=(148, 163, 184))
            set_borders(c, bottom={'val': 'single', 'sz': '4', 'color': 'CBD5E1'}, top={'val': 'single', 'sz': '4', 'color': 'CBD5E1'})
    else:
        add_p(doc, "太棒了！本次卷面書寫嚴謹，無明顯錯別字，請繼續保持高水準之文字正確性！", bold=True, color_rgb=(21, 128, 61))

    # 6. 寫作進階鷹架 (若觸發)
    if scaffold:
        format_heading(doc, "參、 寫作進階鷹架 (篇幅拓展與思想深化引導)", 1)
        scaffold_text = (
            f"1. 【細節擴寫引導】：{scaffold.get('detail_guidance', '試著描述當時對方的具體微表情、環境光影與背景環境音。')}\n"
            f"2. 【五感描摹錦囊】：{scaffold.get('senses_guidance', '注入視覺顏色對比、聽覺節奏、觸覺溫度，避免抽象形容詞。')}\n"
            f"3. 【邏輯與情感深度反思】：{scaffold.get('depth_reflection', '思考事件背後普遍價值：這帶給我看待生命或人際關係何種不可逆的轉變？')}"
        )
        add_callout(doc, scaffold_text, title="🏗️ 寫作進階鷹架卡片", border_color="D97706", bg_color="FEFCE8")

    # 7. 逐段眉批與脈絡診斷
    sec_num_str = "肆" if scaffold else "參"
    format_heading(doc, f"{sec_num_str}、 國文老師逐段批註與脈絡診斷", 1)
    default_secs = [
        ("第一段（起 / 開門破題）", "破題清楚，能點出核心主旨，但引導鋪陳稍急躁，建議加強生活情境之帶入感。"),
        ("第二段（承 / 第一事例展開）", "描寫具有生活細節，能具體呈現人事事物，唯在心理轉折處收尾較倉促。"),
        ("第三段（轉 / 第二事例深入）", "事例與第二段性質略有重複，建議採對比視角，深化『由物入理』之哲思轉化。"),
        ("第四段（合 / 總結回扣昇華）", "有總結意識，但感悟稍流於口號化，結論段必須實質回扣前文事例，形成首尾呼應。")
    ]
    target_secs = section_comments if section_comments else default_secs
    for s_title, s_comment in target_secs:
        add_p(doc, f"• 【{s_title}】：{s_comment}", space_after=2)

    # 8. 段落升級改寫示範
    sec_num_str2 = "伍" if scaffold else "肆"
    format_heading(doc, f"{sec_num_str2}、 突破瓶頸：段落升級改寫示範（衝刺 5~6 級分）", 1)
    orig_p = upgrade_demo.get('orig', "當我停下來的時候，我發現世界其實很美麗。以前我每天都走得很匆忙，沒有好好看路上的風景。現在停下腳步，我感受到陽光很溫暖，微風吹過臉龐很舒服，我學會了不要給自己太大壓力。")
    up_p = upgrade_demo.get('upgrade', "猛然煞住趕赴補習班的腳步，氣喘吁吁地倚在紅磚巷口。褪去耳機裡的喧囂，一陣規律而清脆的鋼鐵敲擊聲悄然鑽入耳際。那是巷口修鞋的老伯，正赤著黝黑的臂膀，在微塵浮動的光暈裡專注穿引麻線。粗糙龜裂的手指一次次沉穩穿透厚硬的皮掌，不慌不忙。那一刻我驀地驚覺，自己日復一日在考卷與分數間的倉皇奔逐，究竟耗損了多少專注的魂魄？唯有停下盲奔，方能讀懂在時光沉澱下的真實生命。")
    tips = upgrade_demo.get('tips', "【改寫要訣點撥】：將原先『陽光溫暖、微風舒服』等抽象空泛心靈雞湯，轉化為具體客觀畫面（修鞋老伯、穿引麻線、光暈微塵），完成『客觀畫面特寫 ➔ 人物專注同理 ➔ 自我心靈省思』之三層昇華！")

    tbl_up = doc.add_table(rows=2, cols=1)
    tbl_up.alignment = WD_TABLE_ALIGNMENT.CENTER; tbl_up.autofit = False
    c_u1 = tbl_up.cell(0, 0); c_u1.width = Inches(7.1)
    set_cell_shd(c_u1, "FEF2F2"); set_cell_pad(c_u1, top=45, bottom=45, left=50, right=50)
    set_borders(c_u1, left={'val': 'single', 'sz': '16', 'color': 'DC2626'})
    p_u1 = c_u1.paragraphs[0]; p_u1.paragraph_format.space_before = Pt(0); p_u1.paragraph_format.space_after = Pt(0)
    r_u1_t = p_u1.add_run("【學生原始段落（4 級分常見通病：抽象抒情多、無客觀細節）】：\n")
    set_font(r_u1_t, name='微軟正黑體', size_pt=8.5, bold=True, color_rgb=(185, 28, 28))
    r_u1_c = p_u1.add_run(orig_p); set_font(r_u1_c, name='微軟正黑體', size_pt=8.0, bold=False, color_rgb=(153, 27, 27))

    c_u2 = tbl_up.cell(1, 0); c_u2.width = Inches(7.1)
    set_cell_shd(c_u2, "EFF6FF"); set_cell_pad(c_u2, top=45, bottom=45, left=50, right=50)
    set_borders(c_u2, left={'val': 'single', 'sz': '16', 'color': '2563EB'})
    p_u2 = c_u2.paragraphs[0]; p_u2.paragraph_format.space_before = Pt(0); p_u2.paragraph_format.space_after = Pt(0)
    r_u2_t = p_u2.add_run("【名師升級示範（5~6 級分典範：注入生活特寫、由物入理）】：\n")
    set_font(r_u2_t, name='微軟正黑體', size_pt=8.5, bold=True, color_rgb=(30, 58, 138))
    r_u2_c = p_u2.add_run(up_p); set_font(r_u2_c, name='微軟正黑體', size_pt=8.0, bold=False, color_rgb=(30, 41, 59))

    add_p(doc, tips, bold=True, color_rgb=(30, 58, 138))

    # 9. 老師完整昇華示範範文 (可選，800~1200字)
    if model_essay:
        sec_num_str3 = "陸" if scaffold else "伍"
        format_heading(doc, f"{sec_num_str3}、 國文老師示範寫作 (完整版昇華範文)", 1)
        tbl_m = doc.add_table(rows=1, cols=1)
        tbl_m.alignment = WD_TABLE_ALIGNMENT.CENTER; tbl_m.autofit = False
        c_m = tbl_m.cell(0, 0); c_m.width = Inches(7.1)
        set_cell_shd(c_m, "F8FAFC"); set_cell_pad(c_m, top=60, bottom=60, left=70, right=70)
        set_borders(c_m, left={'val': 'single', 'sz': '24', 'color': '0284C7'})
        p_m = c_m.paragraphs[0]; p_m.paragraph_format.line_spacing = 1.3
        r_m_title = p_m.add_run(f"〈{topic}〉示範作文\n\n")
        set_font(r_m_title, name='微軟正黑體', size_pt=11.0, bold=True, color_rgb=(2, 132, 199))
        r_m_body = p_m.add_run(model_essay)
        set_font(r_m_body, name='微軟正黑體', size_pt=9.0, bold=False, color_rgb=(30, 41, 59))

def build_single_report(report_data, output_path):
    """產出單一學生專屬診斷單"""
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(0.40)
        s.bottom_margin = Inches(0.40)
        s.left_margin = Inches(0.45)
        s.right_margin = Inches(0.45)

    render_student_section(doc, report_data, is_batch=False)
    doc.save(output_path)
    print(f"單生作文評閱診斷單已成功產出：{output_path}")
    return output_path

def build_batch_class_report(batch_data, output_path):
    """
    產出全班作文評閱彙整手冊 (Class-wide Batch Report)
    包含：
    1. 封面全班成績總覽表
    2. 強制分頁：每位學生獨立成篇
    3. 強制分頁：全班整體教學診斷與課堂精進方針
    """
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = Inches(0.50)
        s.bottom_margin = Inches(0.50)
        s.left_margin = Inches(0.50)
        s.right_margin = Inches(0.50)

    title = batch_data.get('title', "寫作測驗全班成績總覽與個別化評閱手冊")
    subtitle = batch_data.get('subtitle', "國中教育會考規準 ｜ 心測中心量規 + 名師嚴格閱卷指標")
    students = batch_data.get('students', [])
    class_diagnosis = batch_data.get('class_diagnosis', None)

    # 1. 封面抬頭
    p_t = doc.add_paragraph()
    p_t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_t.add_run(title)
    set_font(r_t, name='微軟正黑體', size_pt=20, bold=True, color_rgb=(30, 58, 138))

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(12)
    r_sub = p_sub.add_run(subtitle)
    set_font(r_sub, name='微軟正黑體', size_pt=11, bold=False, color_rgb=(71, 85, 105))

    # 2. 全班成績總覽表
    format_heading(doc, "📊 全班寫作測驗成績與評定總覽表", 1)
    tbl_sum = doc.add_table(rows=1, cols=6)
    tbl_sum.alignment = WD_TABLE_ALIGNMENT.CENTER; tbl_sum.autofit = False
    sum_w = [Inches(1.0), Inches(1.1), Inches(1.8), Inches(1.0), Inches(0.9), Inches(1.7)]
    headers = ["座號/姓名", "學段年級", "寫作題目", "會考級分", "百分制", "核心亮點與硬傷短評"]
    for i, h in enumerate(headers):
        c = tbl_sum.rows[0].cells[i]; c.text = h; c.width = sum_w[i]
        set_cell_shd(c, "1E3A8A"); set_cell_pad(c, top=50, bottom=50, left=40, right=40)
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.runs[0], name='微軟正黑體', size_pt=8.5, bold=True, color_rgb=(255, 255, 255))

    for r_i, s_info in enumerate(students):
        row = tbl_sum.add_row()
        bg = "DCFCE7" if s_info.get('score', 0) >= 6 else ("F8FAFC" if r_i % 2 == 1 else "FFFFFF")
        vals = [
            s_info.get('student_info', f"學生 {r_i+1}"),
            s_info.get('stage', "國中九年級"),
            f"〈{s_info.get('topic', '未定題')}〉",
            f"{s_info.get('score', 4)} 級分",
            f"{s_info.get('percent_score', 80)} 分",
            s_info.get('one_sentence_review', "論述尚稱通順。")[:30] + "..."
        ]
        for c_i, val in enumerate(vals):
            c = row.cells[c_i]; c.text = val; c.width = sum_w[c_i]
            set_cell_shd(c, bg); set_cell_pad(c, top=40, bottom=40, left=40, right=40)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_i in [0, 1, 3, 4] else WD_ALIGN_PARAGRAPH.LEFT
            bold_flag = (c_i in [0, 3])
            f_color = (21, 128, 61) if (c_i == 3 and s_info.get('score', 0)>=6) else (15, 23, 42)
            set_font(p.runs[0], name='微軟正黑體', size_pt=8.0, bold=bold_flag, color_rgb=f_color)
            set_borders(c, bottom={'val': 'single', 'sz': '4', 'color': 'CBD5E1'}, top={'val': 'single', 'sz': '4', 'color': 'CBD5E1'})

    # 封面結束強制換頁
    doc.add_page_break()

    # 3. 逐生渲染章節（每生結束後強制換頁）
    for s_idx, s_data in enumerate(students):
        render_student_section(doc, s_data, is_batch=True)
        if s_idx < len(students) - 1 or class_diagnosis:
            doc.add_page_break()

    # 4. 全班整體教學診斷專頁
    if class_diagnosis:
        format_heading(doc, "🎯 全班整體寫作教學診斷與課堂精進方針", 1)
        add_p(doc, f"主筆名師：{class_diagnosis.get('instructor', '國文領域教學研究會')}", bold=True, color_rgb=(30, 58, 138))
        
        # 普遍盲點
        add_callout(doc, class_diagnosis.get('common_blindspots', '普遍缺乏具體生活客觀畫面，感悟流於心靈雞湯口號化；事例多單一，未能做到一例一段。'),
                    title="🚨 全班普遍失分盲點 (死穴分析)", border_color="DC2626", bg_color="FEF2F2")
        
        # 精進活動
        format_heading(doc, "下一階段課堂增能與對讀研討活動規劃：", 2)
        activities = class_diagnosis.get('activities', [
            "1. 【五感特寫盲寫練習】：播放無聲市集影片 3 分鐘，限時 100 字具體刻畫老者眼神與攤販蒸汽。",
            "2. 【範文對讀解剖會】：以 4 級分與 6 級分同題範文進行紅藍筆標註，找出『由物入理』之關鍵過渡句。",
            "3. 【天頭錯字互批防禦戰】：同儕交換稿紙抓出常犯錯別字與一逗到底語病。"
        ])
        for act in activities:
            add_p(doc, act, space_after=3)

    doc.save(output_path)
    print(f"全班作文評閱彙整手冊已成功產出：{output_path}")
    return output_path

if __name__ == '__main__':
    parser = argparse.ArgumentParser(description="作文專業評閱診斷單/全班彙整手冊 Word 生成工具 (essay_report_builder)")
    parser.add_argument('--topic', default="當我停下來", help="寫作題目")
    parser.add_argument('--student', default="九年級 學生", help="學生座號姓名或年級")
    parser.add_argument('--score', type=int, default=4, help="評定級分 (0~6)")
    parser.add_argument('--tier', default="中卷穩健 (B++)", help="等第類別")
    parser.add_argument('--stage', default="國中九年級", help="學段年級（國小/國中）")
    parser.add_argument('--output', default="會考作文評閱診斷單.docx", help="輸出 Word 檔案路徑")
    parser.add_argument('--data-json', default=None, help="傳入單生或全班評審診斷之 JSON 檔案")
    
    args = parser.parse_args()
    
    if args.data_json and os.path.exists(args.data_json):
        with open(args.data_json, 'r', encoding='utf-8') as f:
            loaded_data = json.load(f)
        if 'students' in loaded_data:
            build_batch_class_report(loaded_data, args.output)
        else:
            build_single_report(loaded_data, args.output)
    else:
        single_data = {
            'topic': args.topic,
            'student_info': args.student,
            'score': args.score,
            'tier': args.tier,
            'stage': args.stage
        }
        build_single_report(single_data, args.output)
